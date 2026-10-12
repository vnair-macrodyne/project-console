"""
console_check_purchasing_omnibus.py — live check of the Procurement Omnibus report and its Excel export.
READ-ONLY.  (2026-10-12)

Runs po_omnibus LIVE for the sample projects, prints timing, the stage counts, and checks that the
Omnibus agrees with the four reports it is built from: To procure rows = Lines to Order rows; draft /
on order / received rows = All Purchase Orders rows; overdue rows = PO Exceptions rows (+ overdue drafts).
Then writes the Excel export to Procurement_Omnibus_check.xlsx in this folder and lists its tabs.

Run:  python console_check_purchasing_omnibus.py              # 260178 250005 220154
      python console_check_purchasing_omnibus.py 250005       # or name projects
Paste the whole output back.
"""
import sys, time, io

PIDS = [int(a) for a in sys.argv[1:] if a.isdigit()] or [260178, 250005, 220154]


def main():
    from console_web.queries import LiveQueryService
    from console_web import exporters
    print(f"console_check_purchasing_omnibus — {time.strftime('%Y-%m-%d %H:%M')}  projects {PIDS}")
    s = LiveQueryService(); t0 = time.perf_counter()
    r = s.run("po_omnibus", PIDS)
    print(f"   built in {time.perf_counter() - t0:.1f}s")
    r.to_dict()
    det = [x for x in r.rows if x.get("_kind") == "detail"]
    print("   cards: " + " · ".join(f"{c.label} {c.value}" for c in r.cards))
    by = {}
    for x in det:
        by[x["FlowStage"]] = by.get(x["FlowStage"], 0) + 1
    print("   stages: " + ", ".join(f"{k} {v:,}" for k, v in by.items()))
    sheets = dict(r.export["sheets"])
    n = lambda res: sum(1 for x in res.rows if x.get("_kind") == "detail")
    print(f"   To procure {by.get('To procure', 0):,} vs Lines to Order {n(sheets['Lines to Order']):,}")
    po_rows = sum(v for k, v in by.items() if k != "To procure")
    print(f"   PO stages {po_rows:,} vs All Purchase Orders {n(sheets['All Purchase Orders']):,}")
    od = sum(1 for x in det if str(x.get("Status") or "").startswith("Overdue"))
    print(f"   Overdue {od:,} vs PO Exceptions {n(sheets['PO Exceptions']):,} (exceptions also count overdue drafts)")
    t1 = time.perf_counter()
    data = exporters.to_xlsx(r)
    open("Procurement_Omnibus_check.xlsx", "wb").write(data)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True)
    print(f"   Excel written in {time.perf_counter() - t1:.1f}s: " + ", ".join(f"{ws.title} ({ws.max_row - 1} rows)" for ws in wb.worksheets))
    print("\nDONE — paste everything above.")


if __name__ == "__main__":
    main()
