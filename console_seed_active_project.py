"""
console_seed_active_project.py — refresh Reporting.tblActiveProject from the Exec Dashboard sheet.

The ACTIVE / open project list is maintained by Helen Young and Mike Plata on the "Worksheet" tab of
"Macrodyne - Executive Dashboard.xlsx". This stages that list into the Console store so Sextant can
scope the Purchasing Exception Report to active projects and reconcile its tracked-project set. Full
refresh (DELETE + reinsert) — a derived snapshot, not a versioned record.

Authoritative active set = the "Sequence" list on the sheet (confirmed with Vijay 2026-10-15: 17
projects incl. GMR/260168). Customer + Priority are enriched from the "Projects #" table below it. If
the sheet's "yes" flags disagree with the Sequence list, the difference is printed (so sheet drift is
visible) but the Sequence list wins.

Run on the app host (reaches the Console store); needs a local copy of the workbook:
    python console_seed_active_project.py --xlsx "Macrodyne - Executive Dashboard.xlsx"
    python console_seed_active_project.py --xlsx "...xlsx" --dry-run   # parse + show, no writes

Idempotent. If the table is missing it is created (matches sql/020_active_project.sql).
"""
from __future__ import annotations

import argparse
import os

import console_config  # noqa: F401 — loads .env into os.environ

DDL = """
IF NOT EXISTS (SELECT 1 FROM sys.tables t
               JOIN sys.schemas s ON s.schema_id = t.schema_id
               WHERE s.name = 'Reporting' AND t.name = 'tblActiveProject')
CREATE TABLE Reporting.tblActiveProject (
    ProjectID   INT           NOT NULL PRIMARY KEY,
    Customer    NVARCHAR(200) NULL,
    Priority    INT           NULL,
    Source      NVARCHAR(100) NOT NULL CONSTRAINT DF_ActiveProject_Source     DEFAULT ('ExecDashboard/Worksheet'),
    ImportedAt  DATETIME      NOT NULL CONSTRAINT DF_ActiveProject_ImportedAt DEFAULT (GETDATE())
);
"""


def _store():
    try:
        from console.infra.connections import console_connection
        return console_connection()
    except Exception:
        import console_store
        return console_store.console_connection()


def parse_sheet(path):
    """Return (sequence_list[int], table{ProjectID: (customer, flag, priority)}) from the Worksheet tab."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Worksheet" not in wb.sheetnames:
        raise RuntimeError(f"No 'Worksheet' tab in {path!r}; tabs = {wb.sheetnames}")
    rows = [list(r) for r in wb["Worksheet"].iter_rows(values_only=True)]

    # 1) the Sequence list: find the 'Sequence' label, then the next cell containing ';'
    seq, anchor = [], None
    for i, r in enumerate(rows):
        if r and r[0] is not None and str(r[0]).strip().lower() == "sequence":
            anchor = i
            break
    if anchor is not None:
        for r in rows[anchor + 1:]:
            if r and r[0] is not None and ";" in str(r[0]):
                seq = [int(x) for x in str(r[0]).split(";") if x.strip().isdigit()]
                break

    # 2) the table under the 'Projects #' header: [ProjectID, Customer, in-extract flag, priority]
    hdr = None
    for i, r in enumerate(rows):
        if r and r[0] is not None and str(r[0]).strip().lower().startswith("project"):
            hdr = i
            break
    table = {}
    if hdr is not None:
        for r in rows[hdr + 1:]:
            if not r or r[0] is None:
                continue
            try:
                pid = int(r[0])
            except (TypeError, ValueError):
                continue
            cust = str(r[1]).strip() if len(r) > 1 and r[1] is not None else None
            flag = str(r[2]).strip().lower() if len(r) > 2 and r[2] is not None else ""
            prio = None
            if len(r) > 3 and r[3] is not None:
                try:
                    prio = int(r[3])
                except (TypeError, ValueError):
                    prio = None
            table[pid] = (cust, flag, prio)
    return seq, table


def build_active(seq, table):
    """Active set = the Sequence list (de-duped, order preserved); enrich from the table; warn on drift."""
    active = list(dict.fromkeys(seq))
    yes = {pid for pid, (c, f, p) in table.items() if f == "yes"}
    only_seq = [p for p in active if p not in yes]
    only_yes = [p for p in yes if p not in active]
    if only_seq or only_yes:
        print(f"NOTE: sheet drift — in Sequence but not 'yes'-flagged: {only_seq}; "
              f"'yes'-flagged but not in Sequence: {only_yes}")
    recs = []
    for pid in active:
        cust, flag, prio = table.get(pid, (None, "", None))
        recs.append((pid, cust, prio))
    return recs


def refresh(store, recs):
    cur = store.cursor()
    try:
        cur.execute(DDL)
        store.commit()
    except Exception as e:
        try:
            store.rollback()
        except Exception:
            pass
        cur.execute("SELECT OBJECT_ID('Reporting.tblActiveProject', 'U')")
        if cur.fetchone()[0] is None:
            raise RuntimeError(
                "Reporting.tblActiveProject does not exist and this login cannot CREATE it "
                f"({type(e).__name__}). Create it once in SSMS with an owner account:\n"
                "    run sql/020_active_project.sql,  then GRANT INSERT, DELETE ON "
                "Reporting.tblActiveProject TO <app login>;\n"
                "then re-run this refresh (which only needs INSERT/DELETE).") from e
        print("note: table exists; skipping CREATE (no CREATE right on this login).")
    cur.execute("DELETE FROM Reporting.tblActiveProject;")
    cur.executemany(
        "INSERT INTO Reporting.tblActiveProject (ProjectID, Customer, Priority) VALUES (?,?,?)",
        recs)
    store.commit()
    return len(recs)


def main():
    ap = argparse.ArgumentParser(description="Refresh Reporting.tblActiveProject from the Exec Dashboard sheet.")
    ap.add_argument("--xlsx", default="Macrodyne - Executive Dashboard.xlsx",
                    help="Path to the Executive Dashboard workbook (has the 'Worksheet' tab).")
    ap.add_argument("--dry-run", action="store_true", help="Parse + show; no DB writes")
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        raise SystemExit(f"Workbook not found: {args.xlsx!r}  (pass --xlsx <path>)")

    seq, table = parse_sheet(args.xlsx)
    recs = build_active(seq, table)
    print(f"active projects: {len(recs)}")
    for pid, cust, prio in recs:
        print(f"   {pid}  {('P'+str(prio)) if prio else '  ':>4}  {cust or ''}")

    if args.dry_run:
        print("\n(dry run — no writes)")
        return

    store = _store()
    try:
        n = refresh(store, recs)
        print(f"\nReporting.tblActiveProject refreshed: {n} rows.")
    finally:
        try:
            store.close()
        except Exception:
            pass


if __name__ == "__main__":
    main()
