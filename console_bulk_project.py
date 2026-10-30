"""
console_bulk_project.py -- bulk ONBOARD / OFFBOARD projects on the Console from a CSV/Excel list.

Reads project numbers (one per row, first column) and applies ONE action to all of them, reusing the
same add_project / remove_project the Plan-page buttons use -- so it's consistent and idempotent:
  * onboard  = clear any exclusion + bank the project's ETO budget (skipped if it has no ETO budget)
  * offboard = add it to the exclusion list (reversible; budget & plan/progress kept)

Run on the app host (reaches the Console + ETO):
    python console_bulk_project.py --file projects.xlsx --action onboard
    python console_bulk_project.py --file projects.csv  --action offboard --dry-run

The file: a single column of project numbers, one per row (a header cell that isn't a number is
skipped). .xlsx/.xlsm (first sheet, first column) or .csv/.tsv/.txt (first column).
"""
import argparse
import csv
import os
import sys

import console_config  # noqa: F401 -- loads .env into os.environ


def _asint(v):
    try:
        if v is None:
            return None
        s = str(v).strip().replace(",", "")
        return int(float(s)) if s else None
    except (ValueError, TypeError):
        return None


def read_ids(path):
    ext = os.path.splitext(path)[1].lower()
    raw = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            if row:
                raw.append(_asint(row[0]))
    elif ext in (".csv", ".tsv", ".txt"):
        delim = "\t" if ext == ".tsv" else ","
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.reader(f, delimiter=delim):
                if row:
                    raw.append(_asint(row[0]))
    else:
        raise SystemExit(f"Unsupported file type {ext!r}; use .xlsx/.xlsm/.csv/.tsv/.txt")
    seen, out = set(), []
    for i in raw:
        if i is not None and i not in seen:
            seen.add(i)
            out.append(i)
    return out


def main():
    ap = argparse.ArgumentParser(description="Bulk onboard/offboard Console projects from a CSV/Excel list.")
    ap.add_argument("--file", required=True, help="CSV/Excel with project numbers (first column).")
    ap.add_argument("--action", required=True, choices=["onboard", "offboard"],
                    help="onboard = add to console (+bank budget); offboard = remove (reversible).")
    ap.add_argument("--entered-by", default="bulk-load", help="Name stamped on the change.")
    ap.add_argument("--dry-run", action="store_true", help="Parse + show the plan; write nothing.")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        raise SystemExit(f"File not found: {args.file}")
    ids = read_ids(args.file)
    print(f"Parsed {len(ids)} project number(s) from {args.file}\n")
    if not ids:
        raise SystemExit("No project numbers found (expected them in the first column).")

    from console_web.pm import make_pm_service
    svc = make_pm_service(demo=False)
    budgeted = set(svc._budgeted_ids())      # banked AND not excluded (currently visible)
    excluded = set(svc._excluded_ids())

    done = already = skipped = errors = 0
    for pid in ids:
        on_console = pid in budgeted
        is_excl = pid in excluded
        if args.action == "onboard":
            if on_console and not is_excl:
                print(f"  {pid}: already on the console"); already += 1; continue
            if args.dry_run:
                print(f"  {pid}: WOULD onboard (currently {'excluded' if is_excl else 'not banked'})"); continue
            try:
                svc.add_project(pid, args.entered_by)
                print(f"  {pid}: ONBOARDED"); done += 1
            except ValueError as e:                       # e.g. no ETO budget to bring in
                print(f"  {pid}: SKIP -- {e}"); skipped += 1
            except Exception as e:
                print(f"  {pid}: ERROR -- {e}"); errors += 1
        else:  # offboard
            if is_excl:
                print(f"  {pid}: already off the console"); already += 1; continue
            if args.dry_run:
                print(f"  {pid}: WOULD offboard"); continue
            try:
                svc.remove_project(pid, args.entered_by)
                print(f"  {pid}: OFFBOARDED"); done += 1
            except Exception as e:
                print(f"  {pid}: ERROR -- {e}"); errors += 1

    tag = "  [DRY-RUN -- nothing written]" if args.dry_run else ""
    print(f"\n{args.action}: {done} changed, {already} already in that state, "
          f"{skipped} skipped (no ETO budget), {errors} error(s){tag}")


if __name__ == "__main__":
    main()
