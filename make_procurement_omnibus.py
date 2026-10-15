"""
make_procurement_omnibus.py — ONE-OFF Procurement Omnibus workbook for Jason (requested via Peter, 12 Oct 2026).
READ-ONLY against ETO. Not part of the console: it reuses the console's purchasing reports so the numbers
match what Sextant shows, and writes one Excel file.

Workbook tabs
  Omnibus                — one row per line through the whole flow, Stage first, then the Procurement
                           Exceptions columns (Peter's 27-column layout):
                             To procure            ETO's Items to Procure (released / requested, no PO yet)
                             Requisition (draft PO) PO raised in ETO, not yet printed or emailed
                             On order              PO sent, not fully received (Status: Open / Overdue / partial)
                             Received              fully received
  Lines to Order         — the console report, as run at the same moment
  Purchase Requisitions  — idem
  PO Exceptions          — idem
  All Purchase Orders    — idem
  Info                   — run time, projects, stage counts, column notes

Run from the repo root where ETO and the Reporting DB are reachable:
    python make_procurement_omnibus.py                      # all projects active on Sextant (budget-current, minus Remove-from-Console)
    python make_procurement_omnibus.py 250005 260178        # or name projects
Writes Procurement_Omnibus_<YYYYMMDD>.xlsx here and prints a summary.
"""
import sys, time, datetime as _dt

STAGES = ["To procure", "Requisition (draft PO)", "On order", "Received"]


def tracked_projects():
    """Projects ACTIVE on Sextant = budget-current MINUS Remove-from-Console exclusions — the exact
    set the console's project selector shows (queries.py LiveQueryService.list_projects)."""
    from console.infra.connections import console_connection
    c = console_connection(); cur = c.cursor()
    cur.execute("SELECT DISTINCT ProjectID FROM Reporting.vw_Console_BudgetCurrent ORDER BY ProjectID")
    ids = [int(r[0]) for r in cur.fetchall()]
    try:
        cur.execute("SELECT ProjectID FROM Reporting.tblConsoleProjectExclusion")
        excl = {int(r[0]) for r in cur.fetchall()}
    except Exception:
        excl = set()
    c.close()
    return [p for p in ids if p not in excl]


def po_order_dates(pids):
    """{PurchaseOrderID (str): order date} for every PO touching the given projects — the PO header
    PurchaseDate (one date per PO), merged onto the Omnibus PO rows by PO number."""
    from console.infra.connections import eto_connection
    ids = ",".join(str(int(p)) for p in pids)
    c = eto_connection(); cur = c.cursor()
    cur.execute("SELECT DISTINCT poh.PurchaseOrderID, CAST(poh.PurchaseDate AS date) "
                "FROM dbo.vwPurchaseOrderHeader poh "
                "JOIN dbo.vwPurchaseOrderDetails pod ON pod.PurchaseOrderID = poh.PurchaseOrderID "
                f"WHERE pod.ProjectID IN ({ids})")
    m = {str(r[0]): r[1] for r in cur.fetchall()}
    c.close()
    return m


def num(v):
    try:
        f = float(v)
        return None if f != f else f
    except (TypeError, ValueError):
        return None


def omnibus_rows(lto, items, cols_exc, podates):
    import pandas as pd
    out = []
    for r in lto.rows:
        if r.get("_kind") != "detail":
            continue
        q = num(r.get("ToOrder"))
        out.append({
            "Stage": "To procure", "Buyer": "", "ReqNo": "", "RFQDate": r.get("RFQDate") or "",
            "PO": "", "OrderDate": "", "Cost": None, "Vendor": "",
            "ProjectID": "" if r.get("ProjectID") is None else str(r.get("ProjectID")),
            "JobName": r.get("JobName") or "", "Code": r.get("Machine") or "", "Item": r.get("ItemNo") or "",
            "Description": r.get("Description") or "", "Category": r.get("Category") or "",
            "DrawingNo": r.get("DrawingNo") or "", "EngRelease": r.get("ReleaseDate") or "",
            "PlannedShip": "", "PlannedReceipt": "", "RevisedReceipt": "", "ReceiptDate": "",
            "Status": "To procure: " + ("" if q is None else (str(int(q)) if q.is_integer() else f"{q:g}")),
            "LastUpdated": "", "MfgBegin": r.get("MfgBegin") or "", "PermitDates": "",
            "LeadTime": num(r.get("LeadTime")), "Oversized": r.get("Oversized") or "",
            "Inspected": r.get("Inspected") or "", "Critical": r.get("Critical") or "", "LLT": r.get("LLT") or "",
        })
    if items is not None and not items.empty:
        for _, r in items.iterrows():
            st = str(r.get("Status") or "")
            stage = "Requisition (draft PO)" if st == "Draft" else ("Received" if st == "Received" else "On order")
            row = {"Stage": stage}
            for k, *_ in cols_exc:
                v = r.get(k)
                row[k] = None if (isinstance(v, float) and pd.isna(v)) else v
            row["OrderDate"] = podates.get(str(r.get("PO")), "")   # PO header PurchaseDate, by PO number
            if stage == "Requisition (draft PO)":
                row["PO"] = ""          # a draft is a requisition: its number shows under Purchasing Requisition #
            out.append(row)
    order = {s: i for i, s in enumerate(STAGES)}
    out.sort(key=lambda x: (str(x.get("ProjectID")), str(x.get("Code")), str(x.get("Item")), order.get(x["Stage"], 9)))
    return out


def write_workbook(path, omni, omni_cols, reports, info_lines):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from console_web import exporters
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    widths = {"text": 20, "date": 12, "money": 14, "money2": 14, "num": 11, "int": 10, "id": 10}

    def table(ws, cols, rows):
        # cols: list of (key, label, type, align)
        for j, (k, lbl, typ, al) in enumerate(cols, 1):
            c = ws.cell(1, j, lbl); c.fill = head_fill; c.font = head_font
            c.alignment = Alignment(horizontal=al, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = max(widths.get(typ, 12), min(len(lbl) + 3, 34))
        for i, row in enumerate(rows, 2):
            for j, (k, lbl, typ, al) in enumerate(cols, 1):
                cell = ws.cell(i, j)
                exporters._put_cell(cell, row.get(k), typ)
                cell.font = body_font
                cell.alignment = Alignment(horizontal=al)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 1)}"
        ws.freeze_panes = "B2"

    wb = Workbook()
    ws = wb.active; ws.title = "Omnibus"
    table(ws, omni_cols, omni)
    for name, res in reports:
        fcols, frows = exporters._flatten(res)
        table(wb.create_sheet(name[:31]), [(c.key, c.label, c.type, c.align) for c in fcols], frows)
    info = wb.create_sheet("Info")
    for i, line in enumerate(info_lines, 1):
        a = info.cell(i, 1, line[0]); a.font = Font(name="Arial", bold=(i == 1 or line[1] is not None), size=12 if i == 1 else 10)
        if line[1] is not None:
            info.cell(i, 2, line[1]).font = body_font
    info.column_dimensions["A"].width = 34; info.column_dimensions["B"].width = 90
    wb.save(path)


def main():
    from console_web.queries import LiveQueryService
    from console_web import etospec
    pids = [int(a) for a in sys.argv[1:] if a.isdigit()] or tracked_projects()
    t0 = time.perf_counter()
    print(f"make_procurement_omnibus — {time.strftime('%Y-%m-%d %H:%M')}  {len(pids)} project(s)")
    svc = LiveQueryService()
    lto = svc.run("released_toorder", pids)
    req = svc.run("po_to_order", pids)
    exc = svc.run("po_exceptions", pids)
    lst = svc.run("po_listing", pids)
    items = (lst.export or {}).get("items")
    podates = po_order_dates(pids)
    omni = omnibus_rows(lto, items, etospec.COLS_EXC, podates)
    omni_cols = [("Stage", "Stage", "text", "left"), ("OrderDate", "Order Date", "date", "left")] + \
        [(k, l, t, a) for (k, l, t, a, w) in etospec.web_columns(etospec.COLS_EXC)]
    det = lambda res: sum(1 for x in res.rows if x.get("_kind") == "detail")
    counts = {s: sum(1 for x in omni if x["Stage"] == s) for s in STAGES}
    overdue = sum(1 for x in omni if str(x.get("Status") or "").startswith("Overdue"))
    cost = sum(num(x.get("Cost")) or 0 for x in omni)
    print(f"   built in {time.perf_counter() - t0:.1f}s")
    for s in STAGES:
        print(f"   {s:24} {counts[s]:>7,}")
    print(f"   overdue (on order)       {overdue:>7,}")
    print(f"   check: To procure {counts['To procure']:,} = Lines to Order {det(lto):,}; "
          f"PO stages {sum(counts[s] for s in STAGES[1:]):,} = All Purchase Orders {det(lst):,}")
    today = _dt.date.today()
    info = [("Procurement Omnibus", None),
            ("Generated", time.strftime("%Y-%m-%d %H:%M")),
            ("Source", "ETO production database (read-only), through the Sextant purchasing reports"),
            ("Scope", "Projects active on Sextant at run time: budgeted in the console and not Removed-from-Console "
                      "(the same set the project selector shows). Override by naming project numbers on the command line."),
            ("Projects", ", ".join(str(p) for p in pids)),
            ("", None)] + [(f"{s} (lines)", f"{counts[s]:,}") for s in STAGES] + [
            ("Overdue (on order)", f"{overdue:,}"),
            ("PO value, CAD (Cost column total)", f"${cost:,.2f}"),
            ("", None),
            ("Stage", "To procure = ETO's Items to Procure: released or requested, no PO yet; Status shows the quantity still to order. "
                      "Requisition (draft PO) = PO raised in ETO but not yet printed or emailed. On order = PO sent, not fully "
                      "received. Received = fully received."),
            ("Purchasing Requisition #", "The PO number as raised in ETO. A draft PO is the requisition; ETO keeps the same number once the PO is sent."),
            ("PO #", "Shown once the PO has been printed or emailed."),
            ("Order Date", "The PO header PurchaseDate — when the PO was placed (for a draft, the date it was raised in ETO). Blank on 'To procure' lines (no PO yet)."),
            ("Cost", "PO line total in CAD. Blank until a PO exists."),
            ("Release Date", "First release-log entry for the project and part; blank for parts bought through a Non-BOM request."),
            ("Mfg Begins", "The machine's Mfg Begin date in ETO, where entered."),
            ("Lead Time", "ETO estimated lead time in days; on overdue PO lines, the part's average historical order-to-receipt days where no estimate exists."),
            ("Planned Ship, Permit Dates", "Blank: rules to be set."),
            ("Oversized, Inspected, Critical, LLT", "Item flags from the ETO item master."),
            ("Left out", "Freight, courier, cartage and brokerage lines; cancelled POs."),
            ("Other tabs", "Lines to Order, Purchase Requisitions, PO Exceptions and All Purchase Orders, as run at the same moment.")]
    path = f"Procurement_Omnibus_{today:%Y%m%d}.xlsx"
    write_workbook(path, omni, omni_cols,
                   [("Lines to Order", lto), ("Purchase Requisitions", req), ("PO Exceptions", exc), ("All Purchase Orders", lst)],
                   info)
    print(f"   wrote {path} in {time.perf_counter() - t0:.1f}s total")
    print("\nDONE — paste everything above.")


if __name__ == "__main__":
    main()
